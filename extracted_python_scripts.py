# ==========================================
# SCRIPT 0 (stream index 547)
# ==========================================

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path

out = Path("/mnt/data/BloodTrace_Project_Presentation.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Palette
NAVY = RGBColor(12, 24, 48)
NAVY2 = RGBColor(20, 39, 70)
TEAL = RGBColor(28, 196, 176)
CYAN = RGBColor(80, 202, 224)
WHITE = RGBColor(245, 248, 252)
MUTED = RGBColor(173, 190, 211)
RED = RGBColor(240, 91, 105)
GOLD = RGBColor(245, 190, 83)
DARK = RGBColor(7, 15, 31)
GREEN = RGBColor(89, 204, 151)

def set_bg(slide, color=NAVY):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color

def add_text(slide, x, y, w, h, text, size=18, color=WHITE, bold=False,
             font="Aptos", align=None, valign=MSO_ANCHOR.TOP, margin=0.06):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        if align is not None:
            p.alignment = align
        p.space_after = Pt(8)
    return box

def add_rich_bullets(slide, x, y, w, h, items, size=17, color=WHITE, bullet_color=TEAL):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(12)
    return box

def add_header(slide, title, subtitle=None, num=None):
    add_text(slide, 0.65, 0.35, 11.9, 0.58, title, 27, WHITE, True)
    # accent bar
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(1.03), Inches(0.72), Inches(0.055))
    shape.fill.solid(); shape.fill.fore_color.rgb = TEAL
    shape.line.fill.background()
    if subtitle:
        add_text(slide, 0.68, 1.16, 12.0, 0.42, subtitle, 12, MUTED)
    if num is not None:
        add_text(slide, 12.3, 7.08, 0.45, 0.2, f"{num:02d}", 9, MUTED, align=PP_ALIGN.RIGHT)

def add_card(slide, x, y, w, h, title, body, accent=TEAL, body_size=15):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = NAVY2
    sh.line.color.rgb = RGBColor(44, 68, 100)
    add_text(slide, x+0.18, y+0.16, w-0.36, 0.35, title, 16, accent, True)
    add_text(slide, x+0.18, y+0.62, w-0.36, h-0.75, body, body_size, WHITE)
    return sh

def add_arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line

# Slide 1
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, DARK)
# Decorative circles
for x,y,d,c in [(9.8,0.5,2.4,NAVY2),(10.8,2.0,1.5,RGBColor(24,70,91)),(8.9,3.7,2.7,NAVY2),(11.1,4.4,1.1,RGBColor(23,86,88))]:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background()
add_text(s,0.85,0.75,7.8,0.35,"ENGINEERING PROJECT • PROOF OF CONCEPT",12,TEAL,True)
add_text(s,0.8,1.55,7.8,1.25,"BLOODTRACE",42,WHITE,True)
add_text(s,0.85,2.82,7.3,0.8,"Optical Screening for Possible Residual Blood Contamination",22,CYAN,True)
add_text(s,0.85,4.05,6.8,0.9,"An experimental approach to highlighting candidate residue on reusable medical instruments before further reprocessing.",16,MUTED)
add_text(s,0.85,6.25,7.4,0.35,"Devansh Lila  |  MNIT Jaipur",13,WHITE,True)
# small motif
for i in range(5):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.2+i*0.5), Inches(3.0+(i%2)*0.45), Inches(0.16), Inches(0.16))
    sh.fill.solid(); sh.fill.fore_color.rgb=TEAL; sh.line.fill.background()

# Slide 2
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"The Problem","Why residual soil on reusable instruments matters",2)
add_card(s,0.7,1.85,3.8,2.0,"Residual material","Blood, tissue and other soil can remain on reusable devices if cleaning is incomplete.",RED,16)
add_card(s,4.75,1.85,3.8,2.0,"Hard-to-inspect areas","Joints, grooves and complex surfaces can make visual inspection challenging.",GOLD,16)
add_card(s,8.8,1.85,3.8,2.0,"Need for screening","A supplementary method could help flag areas for closer manual inspection.",TEAL,16)
add_text(s,0.85,4.45,11.8,1.2,"Project question",18,CYAN,True)
add_text(s,0.85,4.92,11.4,1.05,"Can a compact optical setup and image-analysis pipeline highlight candidate residue regions for human review?",24,WHITE,True)
add_text(s,0.85,6.42,11.6,0.35,"Scope: experimental screening concept — not a replacement for validated reprocessing.",12,MUTED)

# Slide 3
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"Where BloodTrace Fits","Reusable-device reprocessing remains the controlling process",3)
steps=[("01","Point-of-use care","Prevent soil from drying; follow device instructions."),("02","Cleaning","Remove soil and debris using validated procedures."),("03","Inspection","Check the device and identify any remaining visible soil."),("04","Disinfection / sterilization","Use the validated process appropriate to the device.")]
xs=[0.7,3.85,7.0,10.15]
for i,(n,t,b) in enumerate(steps):
    add_card(s,xs[i],2.15,2.5,2.45,t,b,TEAL if i!=3 else CYAN,14)
    add_text(s,xs[i]+0.16,1.75,0.6,0.28,n,12,TEAL,True)
    if i<3: add_arrow(s,xs[i]+2.56,3.3,xs[i+1]-0.12,3.3)
add_text(s,0.85,5.15,11.5,0.85,"Potential role: a supplementary screening step after cleaning, to flag candidate areas for closer inspection.",19,WHITE,True)
add_text(s,0.85,6.25,11.6,0.45,"BloodTrace must not replace cleaning, disinfection, sterilization, or established inspection procedures.",12,MUTED)

# Slide 4
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"Proposed Solution","A compact optical screening proof of concept",4)
add_card(s,0.7,1.85,3.8,2.15,"Illumination","365 nm UV LED and 850 nm NIR LED provide two simulated / experimental illumination modes.",TEAL,15)
add_card(s,4.75,1.85,3.8,2.15,"Image & sensing","Camera captures spatial information; a photodiode provides a single-point intensity reading.",CYAN,15)
add_card(s,8.8,1.85,3.8,2.15,"Analysis","Software compares image features and highlights candidate residue regions for review.",GOLD,15)
# process flow
add_text(s,0.9,4.65,11.5,0.35,"Conceptual workflow",15,TEAL,True)
flow=[("Illuminate","UV / NIR"),("Capture","Camera + sensor"),("Process","Feature analysis"),("Flag","Candidate regions")]
fx=[0.9,4.0,7.1,10.2]
for i,(a,b) in enumerate(flow):
    add_card(s,fx[i],5.15,2.15,1.05,a,b,TEAL,12)
    if i<3: add_arrow(s,fx[i]+2.2,5.68,fx[i+1]-0.12,5.68,CYAN,1.7)

# Slide 5
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"System Architecture","Proposed hardware and software connections",5)
# blocks
add_card(s,0.65,2.0,2.4,1.25,"UV / NIR LEDs","Switched via MOSFETs",TEAL,13)
add_card(s,3.55,2.0,2.4,1.25,"Instrument stage","Sample under test",CYAN,13)
add_card(s,6.45,2.0,2.4,1.25,"USB camera","Image → laptop",GOLD,13)
add_card(s,9.35,2.0,3.2,1.25,"Laptop dashboard","Processing + visualization",TEAL,13)
add_arrow(s,3.08,2.62,3.48,2.62)
add_arrow(s,5.98,2.62,6.38,2.62)
add_arrow(s,8.88,2.62,9.28,2.62)
add_card(s,1.4,4.15,3.1,1.35,"Arduino UNO R4 Minima","Controls illumination modes; sends sensor readings over USB serial.",CYAN,13)
add_card(s,5.1,4.15,3.1,1.35,"Photodiode module","Measures intensity at one point; does not locate stains spatially.",GOLD,13)
add_card(s,8.8,4.15,3.1,1.35,"Power & driver stage","MOSFET switching and current-limited LED circuits.",RED,13)
add_arrow(s,2.95,4.08,2.6,3.32,CYAN,1.5)
add_arrow(s,6.6,4.08,7.5,3.32,GOLD,1.5)
add_arrow(s,10.2,4.08,10.5,3.32,RED,1.5)
add_text(s,0.85,6.25,11.6,0.45,"Camera connects directly to the laptop. Do not drive LEDs directly from Arduino GPIO pins.",12,MUTED)

# Slide 6
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"Working Principle","From optical acquisition to candidate-region map",6)
pipeline=[
    ("1. Select sample","Clean surface, contaminated sample, or control"),
    ("2. Acquire image","Capture visible image; use UV/NIR hardware data when available"),
    ("3. Preprocess","Normalize illumination and isolate the analysis region"),
    ("4. Extract features","Measure intensity, contrast, shape and local differences"),
    ("5. Candidate map","Outline regions that merit manual review"),
]
y=1.85
for i,(t,b) in enumerate(pipeline):
    add_card(s,0.9,y+i*0.88,7.3,0.7,t,b,TEAL if i<4 else CYAN,13)
    if i<4: add_arrow(s,4.55,y+i*0.88+0.72,4.55,y+(i+1)*0.88-0.05,TEAL,1.4)
add_card(s,8.8,2.1,3.55,1.45,"Output","Possible residue detected\nAmbiguous — manual inspection\nNo candidate detected",GOLD,14)
add_card(s,8.8,4.0,3.55,1.45,"Confidence","Algorithmic score from image features; not a validated probability of blood.",CYAN,14)
add_text(s,8.85,5.85,3.4,0.65,"UV/NIR panels must be labelled simulated unless captured by real sensors.",11,MUTED)

# Slide 7 (after removals: original slide 9)
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"Experimental Plan","Evaluate both detection performance and failure modes",7)
add_card(s,0.7,1.8,3.8,2.0,"Sample groups","• Clean surfaces\n• Known, safely prepared surrogate residue\n• Water marks and detergent residue controls",TEAL,14)
add_card(s,4.75,1.8,3.8,2.0,"Test variation","• Different stain sizes and drying conditions\n• Multiple instrument geometries\n• Repeat scans under consistent settings",CYAN,14)
add_card(s,8.8,1.8,3.8,2.0,"Record outcomes","• Candidate region count and area\n• False positives and missed detections\n• Repeatability and uncertainty",GOLD,14)
add_text(s,0.85,4.35,11.7,0.4,"Suggested evaluation metrics",16,TEAL,True)
metrics=[("Sensitivity","Contaminated samples correctly flagged"),("Specificity","Clean/control samples correctly left unflagged"),("False-positive rate","Clean/control samples incorrectly flagged"),("Repeatability","Consistency across repeated scans")]
for i,(t,b) in enumerate(metrics):
    x=0.8+(i%2)*6.1; y=4.85+(i//2)*0.85
    add_text(s,x,y,2.05,0.3,t,13,CYAN,True)
    add_text(s,x+2.0,y,3.9,0.55,b,12,WHITE)
add_text(s,0.85,6.75,11.7,0.25,"Keep threshold-tuning images separate from the evaluation set.",10,MUTED)

# Slide 8 (original slide 11)
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"Future Scope","From simulation to a defensible prototype",8)
add_card(s,0.7,1.85,3.8,2.05,"1. Validate optics","Capture actual images under controlled 365 nm UV and 850 nm NIR illumination; document camera sensitivity and geometry.",TEAL,14)
add_card(s,4.75,1.85,3.8,2.05,"2. Improve robustness","Test different surfaces, lighting, residue levels, and confounding materials; quantify uncertainty.",CYAN,14)
add_card(s,8.8,1.85,3.8,2.05,"3. Establish ground truth","Compare candidate maps against independently labelled samples and an appropriate reference method.",GOLD,14)
add_text(s,0.85,4.55,11.6,0.4,"Longer-term opportunities",16,TEAL,True)
add_rich_bullets(s,0.9,5.0,11.5,1.45,[
    "Build a repeatable enclosed optical fixture with controlled illumination.",
    "Create a labelled dataset and report performance with uncertainty intervals.",
    "Consult the institution’s IP cell before public disclosure if patent protection is being considered."
],14)

# Slide 9 References
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); add_header(s,"References","Background sources for reprocessing and infection-control context",9)
refs=[
    ("U.S. Food and Drug Administration","How are reusable medical devices reprocessed?","https://www.fda.gov/medical-devices/reprocessing-reusable-medical-devices/how-are-reusable-medical-devices-reprocessed"),
    ("Centers for Disease Control and Prevention","Disinfection of healthcare equipment","https://www.cdc.gov/infection-control/hcp/disinfection-sterilization/healthcare-equipment.html"),
    ("PubMed Central","Research article (PMCID: PMC7134755)","https://pmc.ncbi.nlm.nih.gov/articles/PMC7134755/"),
]
for i,(org,title,url) in enumerate(refs):
    y=1.85+i*1.35
    add_text(s,0.85,y,11.5,0.32,org,16,TEAL,True)
    add_text(s,0.85,y+0.35,11.5,0.32,title,14,WHITE)
    add_text(s,0.85,y+0.72,11.5,0.35,url,10,CYAN)
add_text(s,0.85,6.15,11.6,0.75,"Interpretation: these sources support the motivation and safety framing; they do not validate BloodTrace’s detection accuracy.",12,MUTED)

# Apply consistent metadata
prs.core_properties.title = "BloodTrace — Optical Screening Proof of Concept"
prs.core_properties.subject = "Engineering project presentation"
prs.core_properties.author = "Devansh Lila, MNIT Jaipur"
prs.core_properties.keywords = "BloodTrace, reusable medical devices, optical screening, image analysis"

prs.save(out)
print(f"Created {out}")


# ==========================================
# SCRIPT 1 (stream index 789)
# ==========================================

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

base = Path("/mnt/data")
source = base / "BloodTrace_College_Project_Proposal.docx"
output = base / "BloodTrace_College_Project_Proposal_Research_Updated.docx"

if not source.exists():
    print("Original proposal file is not available at the expected path.")
    print("Available files:", [p.name for p in base.glob("*")][:80])
else:
    doc = Document(str(source))

    doc.add_page_break()
    h = doc.add_heading("Research Context and Design Implications", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph(
        "The following discussion connects guidance on reusable medical-device reprocessing "
        "to the BloodTrace proof-of-concept. These sources establish the project context and "
        "safety requirements; they do not validate the proposed optical system or establish its accuracy."
    )

    doc.add_heading("1. Relevance of reusable-device reprocessing guidance", level=2)
    doc.add_paragraph(
        "The U.S. Food and Drug Administration describes reprocessing as a sequence of steps "
        "intended to prepare reusable medical devices for reuse. Cleaning removes soil and debris "
        "and is an important prerequisite to appropriate disinfection or sterilization. BloodTrace "
        "could be explored as an experimental screening aid for possible residual surface material, "
        "but it must not replace validated cleaning, disinfection, sterilization, or inspection procedures."
    )

    doc.add_heading("2. Implications from infection-control guidance", level=2)
    doc.add_paragraph(
        "CDC guidance emphasizes the importance of cleaning reusable healthcare equipment before "
        "disinfection or sterilization. A visually clean surface does not demonstrate that microorganisms "
        "are absent. Blood residue and microbial contamination are different targets: detecting or failing "
        "to detect a stain cannot establish whether an instrument is disinfected or sterile."
    )

    doc.add_heading("3. Engineering requirements for BloodTrace", level=2)
    requirements = [
        "Frame the output as a candidate-residue screening result, not a declaration that an instrument is safe or sterile.",
        "Include samples with different stain sizes, drying conditions, instrument geometries, joints, and grooves.",
        "Use negative controls such as clean surfaces, water marks, and detergent residue to assess false positives.",
        "Show an “Ambiguous — Manual Inspection Required” result when image evidence is inconsistent.",
        "Use ordinary photographs to test software only; capture actual wavelength-specific images to evaluate the physical UV/NIR hardware.",
        "Keep threshold-tuning images separate from evaluation images to reduce overly optimistic performance estimates."
    ]
    for item in requirements:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("4. Proposed evaluation metrics", level=2)
    doc.add_paragraph(
        "Evaluate the system against independently established sample labels and, where appropriate, "
        "a suitable reference method. Report sensitivity (the fraction of contaminated samples correctly "
        "flagged), specificity (the fraction of clean/control samples correctly left unflagged), false-positive "
        "rate, false-negative rate, and repeatability across repeated scans. Do not report performance "
        "metrics until sufficient testing has been completed."
    )

    doc.add_heading("5. Limits of the current evidence", level=2)
    doc.add_paragraph(
        "The FDA and CDC sources support the motivation and safety framing of the project, but they do "
        "not validate BloodTrace's proposed 365 nm UV or 850 nm NIR detection approach. Ordinary visible-light "
        "photographs cannot establish how the physical device responds at those wavelengths. UV fluorescence "
        "and NIR reflectance are not, by themselves, specific proof of blood. The project should remain a "
        "proof-of-concept until it has been validated against suitable samples and reference measurements."
    )
    doc.add_paragraph(
        "The third supplied link (https://pmc.ncbi.nlm.nih.gov/articles/PMC7134755/) is included in the "
        "reference list below. Its specific findings have not been summarized here because the full article "
        "text was not available for verification during preparation."
    )

    doc.add_page_break()
    doc.add_heading("Additional References", level=1)

    references = [
        ("U.S. Food and Drug Administration. (n.d.). ", "How are reusable medical devices reprocessed? ",
         "https://www.fda.gov/medical-devices/reprocessing-reusable-medical-devices/how-are-reusable-medical-devices-reprocessed"),
        ("Centers for Disease Control and Prevention. (n.d.). ", "Disinfection of healthcare equipment. ",
         "https://www.cdc.gov/infection-control/hcp/disinfection-sterilization/healthcare-equipment.html"),
        ("PubMed Central. (n.d.). ", "Article available at PubMed Central (PMCID: PMC7134755). ",
         "https://pmc.ncbi.nlm.nih.gov/articles/PMC7134755/"),
    ]
    for author, title, url in references:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(8)
        p.add_run(author)
        run = p.add_run(title)
        run.italic = True
        p.add_run(url)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    r = note.add_run(
        "Safety note: BloodTrace is an experimental screening concept. It must not replace validated "
        "medical-device reprocessing procedures or be used to declare an instrument clean, disinfected, or sterile."
    )
    r.bold = True

    doc.save(str(output))
    print(f"Created updated report: {output}")


# ==========================================
# SCRIPT 2 (stream index 1778)
# ==========================================

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

src = Path("/mnt/data/BloodTrace_College_Project_Proposal.docx")
out = Path("/mnt/data/BloodTrace_College_Project_Proposal_Updated_References.docx")

if not src.exists():
    print("Original document not found at expected path.")
    print("Files in /mnt/data:", [p.name for p in Path("/mnt/data").glob("*")][:50])
else:
    doc = Document(str(src))
    doc.add_page_break()
    heading = doc.add_heading("Additional References", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(
        "The following authoritative guidance and review literature provide background on reusable "
        "medical-device reprocessing, cleaning, disinfection, and sterilization. These sources support "
        "the project context and safety limitations; they do not validate BloodTrace as a blood detector."
    )
    p.paragraph_format.space_after = Pt(10)

    refs = [
        ("U.S. Food and Drug Administration. (n.d.). ", "How are reusable medical devices reprocessed?",
         " https://www.fda.gov/medical-devices/reprocessing-reusable-medical-devices/how-are-reusable-medical-devices-reprocessed"),
        ("Centers for Disease Control and Prevention. (2023, November 28). ", "Disinfection of healthcare equipment.",
         " https://www.cdc.gov/infection-control/hcp/disinfection-sterilization/healthcare-equipment.html"),
        ("Rutala, W. A., & Weber, D. J. (2016). ",
         "Disinfection and sterilization in health care facilities: An overview and current issues.",
         " Infectious Disease Clinics of North America, 30(3), 609–637. https://doi.org/10.1016/j.idc.2016.04.002"),
    ]
    for lead, title, tail in refs:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(8)
        p.add_run(lead)
        r = p.add_run(title)
        r.italic = True
        p.add_run(tail)

    doc.add_paragraph(
        "Access note: Web references checked on 18 October 2026. The FDA describes general steps for "
        "reusable-device reprocessing; CDC guidance emphasizes that appropriate cleaning must precede "
        "disinfection or sterilization. Rutala and Weber review healthcare disinfection and sterilization "
        "principles. BloodTrace remains an experimental screening concept and must not replace validated "
        "reprocessing procedures or be used to declare an instrument sterile."
    )
    doc.save(str(out))
    print(f"Created: {out}")


# ==========================================
# SCRIPT 3 (stream index 1818)
# ==========================================

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

src = Path("/mnt/data/BloodTrace_College_Project_Proposal.docx")
out = Path("/mnt/data/BloodTrace_College_Project_Proposal_Updated_References.docx")

if not src.exists():
    raise FileNotFoundError(f"Original proposal not found at {src}")

doc = Document(str(src))

doc.add_page_break()
heading = doc.add_heading("Additional References", level=1)
heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

p = doc.add_paragraph(
    "The following authoritative guidance and review literature provide background on reusable "
    "medical-device reprocessing, cleaning, disinfection, and sterilization. They support the "
    "project context and safety limitations; they do not validate BloodTrace as a blood detector."
)
p.paragraph_format.space_after = Pt(10)

refs = [
    (
        "U.S. Food and Drug Administration. (n.d.). ",
        "How are reusable medical devices reprocessed?",
        " https://www.fda.gov/medical-devices/reprocessing-reusable-medical-devices/how-are-reusable-medical-devices-reprocessed"
    ),
    (
        "Centers for Disease Control and Prevention. (2023, November 28). ",
        "Disinfection of healthcare equipment.",
        " https://www.cdc.gov/infection-control/hcp/disinfection-sterilization/healthcare-equipment.html"
    ),
    (
        "Rutala, W. A., & Weber, D. J. (2016). ",
        "Disinfection and sterilization in health care facilities: An overview and current issues.",
        " Infectious Disease Clinics of North America, 30(3), 609–637. https://doi.org/10.1016/j.idc.2016.04.002"
    ),
]

for lead, title, tail in refs:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(8)
    p.add_run(lead)
    r = p.add_run(title)
    r.italic = True
    p.add_run(tail)

doc.add_paragraph(
    "Access note: Web references checked on 18 October 2026. The FDA describes general steps "
    "for reusable-device reprocessing; CDC guidance emphasizes that appropriate cleaning must "
    "precede disinfection or sterilization. Rutala and Weber review healthcare disinfection and "
    "sterilization principles. BloodTrace remains an experimental screening concept and must not "
    "replace validated reprocessing procedures or be used to declare an instrument sterile."
)

doc.save(str(out))
print(f"Created: {out}")


# ==========================================
# SCRIPT 4 (stream index 2520)
# ==========================================

import cv2, os
p="/mnt/data/BloodTrace_Simulation_Showcase_1080p.mp4"
cap=cv2.VideoCapture(p)
print(os.path.exists(p), cap.isOpened(), cap.get(cv2.CAP_PROP_FRAME_COUNT), cap.get(cv2.CAP_PROP_FPS), cap.get(cv2.CAP_PROP_FRAME_WIDTH), cap.get(cv2.CAP_PROP_FRAME_HEIGHT))


# ==========================================
# SCRIPT 5 (stream index 2583)
# ==========================================

import cv2, os
p="/mnt/data/BloodTrace_Simulation_Showcase_1080p.mp4"
cap=cv2.VideoCapture(p)
cap.isOpened(), cap.get(cv2.CAP_PROP_FRAME_COUNT), cap.get(cv2.CAP_PROP_FPS), cap.get(cv2.CAP_PROP_FRAME_WIDTH), cap.get(cv2.CAP_PROP_FRAME_HEIGHT)


# ==========================================
# SCRIPT 6 (stream index 3921)
# ==========================================

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch, Circle
import os, textwrap

out_dir = "/mnt/data"
docx_path = os.path.join(out_dir, "BloodTrace_College_Project_Proposal.docx")
diagram_path = os.path.join(out_dir, "BloodTrace_System_Block_Diagram.png")
circuit_path = os.path.join(out_dir, "BloodTrace_Simplified_Circuit.png")

# System block diagram
fig, ax = plt.subplots(figsize=(11, 4.4))
ax.set_xlim(0, 12); ax.set_ylim(0, 5); ax.axis("off")
def box(x,y,w,h,title,sub="",fc="#eef4fb"):
    p = FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.04,rounding_size=0.12",
                       linewidth=1.4,edgecolor="#23415f",facecolor=fc)
    ax.add_patch(p)
    ax.text(x+w/2,y+h*0.62,title,ha="center",va="center",fontsize=10,fontweight="bold")
    if sub: ax.text(x+w/2,y+h*0.30,sub,ha="center",va="center",fontsize=8)
def arrow(x1,y1,x2,y2,label=None):
    ax.annotate("",xy=(x2,y2),xytext=(x1,y1),arrowprops=dict(arrowstyle="->",lw=1.6,color="#23415f"))
    if label: ax.text((x1+x2)/2,(y1+y2)/2+0.16,label,ha="center",fontsize=8)
box(.2,2.0,2.2,1.5,"Instrument","cleaned test surface")
box(2.9,2.0,2.4,1.5,"Optical head","365 nm UV + 850 nm IR\nUSB camera", "#f7f0ff")
box(5.8,2.0,2.0,1.5,"Arduino UNO R4","LED control + sensor\nreading", "#e9f8f0")
box(8.4,2.0,1.4,1.5,"Laptop","USB camera +\nserial data", "#fff5e6")
box(10.2,2.0,1.6,1.5,"Display","image + exploratory\nresidue map", "#fceeee")
arrow(2.4,2.75,2.9,2.75)
arrow(5.3,2.75,5.8,2.75,"control")
arrow(7.8,2.75,8.4,2.75,"USB serial")
arrow(9.8,2.75,10.2,2.75,"OpenCV")
ax.text(6,4.45,"BLOODTRACE — PROPOSED COLLEGE-SCALE PROTOTYPE",ha="center",fontsize=14,fontweight="bold",color="#16324f")
ax.text(6,.8,"The camera supplies spatial image data; the photodiode supplies a separate light-intensity measurement.",ha="center",fontsize=9)
fig.tight_layout()
fig.savefig(diagram_path,dpi=200,bbox_inches="tight")
plt.close(fig)

# Simplified circuit diagram
fig, ax = plt.subplots(figsize=(10, 6))
ax.set_xlim(0, 12); ax.set_ylim(0, 8); ax.axis("off")
ax.text(6,7.6,"SIMPLIFIED ELECTRICAL CONNECTIONS (NOT TO SCALE)",ha="center",fontsize=13,fontweight="bold")
# Arduino block
p=FancyBboxPatch((4.4,2.3),3.2,3.7,boxstyle="round,pad=0.08",linewidth=1.5,edgecolor="#23415f",facecolor="#eef4fb")
ax.add_patch(p); ax.text(6,5.55,"Arduino UNO R4 Minima",ha="center",fontsize=11,fontweight="bold")
for y, t in [(4.85,"D9 → UV MOSFET gate via 1 kΩ"),(4.15,"D10 → IR MOSFET gate via 1 kΩ"),(3.45,"A0 ← photodiode amplifier output"),(2.75,"D2 ← push button to GND")]:
    ax.text(6,y,t,ha="center",fontsize=8.8)
# UV branch
ax.text(.5,6.3,"UV CHANNEL",fontsize=10,fontweight="bold",color="#6b3fa0")
ax.text(.5,5.7,"+5 V → 220 Ω* → UV LED anode",fontsize=9)
ax.text(.5,5.25,"UV LED cathode → Q1 drain",fontsize=9)
ax.text(.5,4.8,"Q1 source → GND; Q1 gate → D9",fontsize=9)
ax.text(.5,4.35,"10 kΩ gate pulldown to GND",fontsize=9)
# IR branch
ax.text(8.2,6.3,"IR CHANNEL",fontsize=10,fontweight="bold",color="#a34a20")
ax.text(8.2,5.7,"+5 V → 220 Ω* → IR LED anode",fontsize=9)
ax.text(8.2,5.25,"IR LED cathode → Q2 drain",fontsize=9)
ax.text(8.2,4.8,"Q2 source → GND; Q2 gate → D10",fontsize=9)
ax.text(8.2,4.35,"10 kΩ gate pulldown to GND",fontsize=9)
# photodiode block
boxp=FancyBboxPatch((8.1,1.1),3.4,2.0,boxstyle="round,pad=0.06",linewidth=1.2,edgecolor="#23415f",facecolor="#f7f7f7")
ax.add_patch(boxp)
ax.text(9.8,2.65,"BPW34 photodiode",ha="center",fontsize=10,fontweight="bold")
ax.text(9.8,2.2,"→ photodiode amplifier module",ha="center",fontsize=8.5)
ax.text(9.8,1.75,"module OUT → A0; VCC → 5 V",ha="center",fontsize=8.5)
ax.text(9.8,1.4,"GND → common GND",ha="center",fontsize=8.5)
# USB
ax.annotate("",xy=(6,1.0),xytext=(6,2.25),arrowprops=dict(arrowstyle="<->",lw=1.4,color="#23415f"))
ax.text(6.2,1.55,"USB-C to laptop\n(power + serial)",fontsize=8,va="center")
ax.text(6,.35,"*220 Ω is an initial value for a single low-current LED only; verify LED forward voltage/current and resistor power before powering.",ha="center",fontsize=8)
fig.tight_layout()
fig.savefig(circuit_path,dpi=200,bbox_inches="tight")
plt.close(fig)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(5)
for s in ["Title","Heading 1","Heading 2","Heading 3"]:
    styles[s].font.name = "Arial"
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.color.rgb = RGBColor(24, 61, 91)
styles["Heading 2"].font.size = Pt(12)
styles["Heading 2"].font.color.rgb = RGBColor(30, 91, 112)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text)); r.bold = bold; r.font.name = "Arial"; r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True); shade(t.rows[0].cells[i], "D9EAF4")
    for row in rows:
        cells = t.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i], v)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t
def bullet(txt):
    doc.add_paragraph(txt, style="List Bullet")
def numbered(txt):
    doc.add_paragraph(txt, style="List Number")

# Cover
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\n")
r=p.add_run("BLOODTRACE"); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor(24,61,91)
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("A Dual-Wavelength Optical Screening Prototype for Possible Residual Blood Contamination on Reusable Medical Instruments")
r.bold=True; r.font.size=Pt(15)
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\nCollege-Level Project Proposal\n\n")
p.add_run("Submitted to: ______________________________\n")
p.add_run("Department: ________________________________\n")
p.add_run("Submitted by: Devansh Lila\n")
p.add_run("Institution: MNIT Jaipur\n")
p.add_run("Date: __________________")
doc.add_page_break()

doc.add_heading("Abstract", level=1)
doc.add_paragraph(
    "This report proposes BloodTrace, a small-scale optical screening prototype intended to help visualize possible residual contamination on reusable medical instruments after cleaning and before sterilization. The design combines 365 nm ultraviolet (UV) illumination, near-infrared (NIR) illumination at approximately 850 nm, a camera, an optional BPW34 photodiode channel, an Arduino UNO R4 Minima, and laptop-based image processing using Python and OpenCV. The Arduino controls illumination and streams sensor readings, while the laptop captures images, compares illumination modes, and overlays regions that differ from a reference image. The output is an exploratory residue map, not a definitive blood identification result. The proposed project is suitable for an educational proof of concept; further controlled experiments would be required to determine sensitivity, specificity, repeatability, and suitability for any clinical use."
)
doc.add_heading("Project status and intended use", level=2)
doc.add_paragraph(
    "Status: proposed prototype; performance has not yet been experimentally established. Intended use: educational research and bench testing on non-clinical test samples. It must not be used to release medical instruments for patient use or to replace validated cleaning, disinfection, sterilization, or monitoring procedures."
)

doc.add_heading("1. Introduction and problem statement", level=1)
doc.add_paragraph(
    "Reusable medical instruments must be thoroughly cleaned before high-level disinfection or sterilization. Residual organic material, including blood and tissue, can interfere with subsequent processing. Visual inspection alone may not reveal small or dried residues, especially around joints, serrations, and complex surfaces. A low-cost screening aid that captures an image and highlights unusual optical regions could support training and research into cleaning verification."
)
doc.add_paragraph(
    "The proposed problem is: Can a compact, low-cost optical system combine UV and NIR illumination with image processing to highlight regions that may contain residual blood or other contamination on a reusable instrument?"
)

doc.add_heading("2. Aim and objectives", level=1)
for x in [
    "Build a compact optical head using UV and NIR LEDs and a camera.",
    "Use an Arduino UNO R4 Minima to switch illumination modes and read an auxiliary optical sensor.",
    "Capture visible/UV/NIR-related images on a laptop and display an exploratory residue map.",
    "Compare clean reference surfaces with controlled test samples and record repeatability.",
    "Document limitations and identify what further validation would be needed before any real-world use."
]: bullet(x)

doc.add_heading("3. Proposed design and working principle", level=1)
doc.add_picture(diagram_path, width=Inches(6.8))
p=doc.add_paragraph("Figure 1. System-level block diagram of the proposed prototype.")
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(
    "The instrument is placed inside a small, matte-black enclosure to reduce ambient light. The system captures a reference image under visible illumination, then captures images with UV and NIR illumination separately. A camera with suitable sensitivity and filtering is needed to obtain meaningful images in the relevant wavelength bands. The optional photodiode records a single-point light-intensity response; it does not independently locate a stain. The laptop aligns and compares images and marks candidate regions for review."
)
doc.add_heading("Why use both UV and IR?", level=2)
doc.add_paragraph(
    "UV illumination may reveal fluorescence from some residues, but fluorescence is not unique to blood and can vary with wavelength, surface finish, cleaning agents, and residue age. NIR illumination can provide a complementary reflectance/absorption contrast channel, but a basic IR LED and camera do not automatically identify hemoglobin. The value of the dual-wavelength approach is therefore a testable hypothesis: combining channels may improve contrast over a single image. This must be assessed experimentally rather than assumed."
)

doc.add_heading("4. Hardware and components", level=1)
table(["Component","Qty.","Purpose / notes"],[
    ("Arduino UNO R4 Minima","1","Controls UV/IR drivers and reads analog sensor data; USB serial to laptop."),
    ("365 nm UV LED or low-current module","1","UV illumination; enclosed to prevent eye/skin exposure."),
    ("850 nm IR LED","1–3","NIR illumination; each LED needs an appropriately sized resistor or current driver."),
    ("USB webcam","1","Image acquisition; test whether it is sensitive to the chosen wavelength. A camera without an IR-cut filter may be needed for NIR."),
    ("BPW34 photodiode","1","Auxiliary optical measurement; use with a suitable amplifier or photodiode amplifier module."),
    ("Photodiode amplifier module","1","Converts small photodiode current to a measurable voltage; check output range and Arduino compatibility."),
    ("AO3400 logic-level N-MOSFETs","2","Low-side switching for LED channels; verify package pinout and current rating."),
    ("1 kΩ resistors","2","Series resistors between Arduino control pins and MOSFET gates."),
    ("10 kΩ resistors","2","Gate pulldowns to keep MOSFETs off during reset."),
    ("220 Ω resistors","1 per LED branch, initially","Starting value only for a single low-current LED; calculate from LED forward voltage and desired current."),
    ("Push button + jumper wires","1 set","Mode change / start trigger and wiring."),
    ("Black enclosure / foam baffles","1","Reduces ambient light and separates illumination from camera view."),
    ("USB-C cable + laptop","1 each","Board power/programming/serial; camera connects separately to laptop USB."),
    ("Breadboard or prototype PCB","1","Assembly for a supervised college prototype.")
])
doc.add_paragraph(
    "Budget note: prices vary by supplier and camera/module quality. The Arduino and camera are likely to dominate cost if not already available in the lab. Prepare a local bill of materials and confirm part availability before procurement."
)

doc.add_heading("5. Electrical connections", level=1)
doc.add_picture(circuit_path, width=Inches(6.8))
p=doc.add_paragraph("Figure 2. Simplified connection overview; verify the exact module and component datasheets before assembly.")
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
table(["UNO R4 Minima pin","Connection"],[
    ("D9 (PWM)","UV-channel MOSFET gate through 1 kΩ; add 10 kΩ gate-to-GND pulldown."),
    ("D10 (PWM)","IR-channel MOSFET gate through 1 kΩ; add 10 kΩ gate-to-GND pulldown."),
    ("A0","Output from photodiode amplifier module, kept within the Arduino analog-input voltage range."),
    ("D2","Push button from D2 to GND; configure as INPUT_PULLUP."),
    ("5V","Supply for compatible low-current sensor amplifier/module, if its specification permits."),
    ("GND","Common ground for Arduino, MOSFET sources, amplifier module, and external LED supply."),
    ("USB-C","Arduino to laptop for power, programming, and serial data."),
    ("USB camera","Connect directly to laptop, not to the Arduino.")
])
doc.add_heading("LED driver wiring", level=2)
doc.add_paragraph(
    "For each LED channel, connect +5 V → current-limiting resistor → LED anode; LED cathode → MOSFET drain; MOSFET source → GND. Connect Arduino D9 or D10 to the MOSFET gate through 1 kΩ, and connect a 10 kΩ pulldown from gate to GND. This is a low-side switching arrangement. Do not power an LED array directly from an Arduino GPIO pin. The 220 Ω value is only an initial example for a single low-current LED; calculate R = (V_supply − V_forward) / I_target for the actual LED and confirm resistor power dissipation."
)
doc.add_heading("Photodiode channel", level=2)
doc.add_paragraph(
    "Use a BPW34 with a compatible photodiode amplifier/transimpedance module for the first prototype rather than assembling an untested amplifier from a sketch. Connect the module output to A0, its supply to the specified rail, and its ground to common GND. Confirm that the module output never exceeds the UNO R4 Minima analog-input range. A photodiode is a single-point sensor and should be treated as a secondary measurement, not the source of the stain map."
)

doc.add_heading("6. Software architecture", level=1)
doc.add_paragraph(
    "Arduino firmware cycles through UV-only, IR-only, and combined modes; reads the auxiliary sensor; and sends mode and reading values over USB serial. Laptop software captures webcam frames, stores reference images, applies background/illumination correction, compares frames, and overlays regions with unusual intensity or texture. Thresholds should initially be adjustable and treated as exploratory, not as validated blood-detection thresholds."
)
doc.add_heading("6.1 Arduino firmware (prototype)", level=2)
arduino_code = r'''const int UV_PIN = 9;
const int IR_PIN = 10;
const int SENSOR_PIN = A0;
const int BUTTON_PIN = 2;

int mode = 0; // 0 = UV, 1 = IR, 2 = UV + IR
bool lastButton = HIGH;
unsigned long lastDebounce = 0;

void setMode() {
  digitalWrite(UV_PIN, mode == 0 || mode == 2 ? HIGH : LOW);
  digitalWrite(IR_PIN, mode == 1 || mode == 2 ? HIGH : LOW);
}

void setup() {
  Serial.begin(115200);
  pinMode(UV_PIN, OUTPUT);
  pinMode(IR_PIN, OUTPUT);
  pinMode(BUTTON_PIN, INPUT_PULLUP);
  digitalWrite(UV_PIN, LOW);
  digitalWrite(IR_PIN, LOW);
  Serial.println("BloodTrace prototype ready");
}

void loop() {
  bool buttonNow = digitalRead(BUTTON_PIN);

  if (lastButton == HIGH && buttonNow == LOW &&
      millis() - lastDebounce > 250) {
    mode = (mode + 1) % 3;
    setMode();
    lastDebounce = millis();
  }
  lastButton = buttonNow;

  int sensorValue = analogRead(SENSOR_PIN);
  Serial.print("MODE=");
  Serial.print(mode);
  Serial.print(",SENSOR=");
  Serial.println(sensorValue);
  delay(100);
}'''
p=doc.add_paragraph()
r=p.add_run(arduino_code); r.font.name="Consolas"; r.font.size=Pt(8)

doc.add_heading("6.2 Laptop image-processing outline", level=2)
for x in [
    "Connect the webcam and capture an image with the LEDs off, then UV-only and IR-only.",
    "Keep camera position, exposure, distance, and instrument orientation fixed.",
    "Correct for background and illumination differences; compare each test frame with a clean reference.",
    "Use conservative thresholding and contour detection to outline candidate regions.",
    "Display the original image, UV frame, NIR frame, and candidate-region overlay side by side.",
    "Label results as 'candidate residue region' or 'possible contamination'; do not label them 'blood confirmed' without validated testing."
]: numbered(x)
doc.add_paragraph(
    "A basic OpenCV thresholding demonstration can highlight differences but cannot, by itself, identify blood chemically. More advanced classification should only be developed after collecting a properly labelled dataset containing clean surfaces, blood-containing test samples, cleaning-agent residues, and other likely interferents."
)

doc.add_heading("7. Experimental plan", level=1)
doc.add_paragraph("Use non-clinical test coupons or retired training instruments only. Do not use the prototype to decide whether patient-care instruments are safe.")
table(["Stage","Method","Record"],[
    ("1. Baseline","Capture repeated images of a clean metal coupon under visible, UV, and NIR modes.","Image variation, background noise, repeatability."),
    ("2. Positive-control feasibility","Under supervisor-approved lab procedures, use a safe, controlled test material or commercially available test standard; avoid handling human blood unless an approved biosafety protocol is in place.","Signal/contrast change versus baseline."),
    ("3. Interference testing","Test common confounders such as water, detergent residue, and surface marks on separate coupons.","False alarms and missed candidate regions."),
    ("4. Repeatability","Repeat each condition at fixed and slightly varied distances/orientations.","Within-condition variation."),
    ("5. Image-map evaluation","Have a reviewer compare marked regions with known sample locations.","Region overlap, false-positive rate, false-negative rate."),
    ("6. Iteration","Adjust enclosure, illumination geometry, exposure, filters, and thresholds.","Versioned settings and performance change.")
])
doc.add_heading("8. Evaluation metrics", level=1)
for x in [
    "Detection sensitivity: fraction of known positive test areas marked by the system.",
    "Specificity: fraction of known clean/interference areas not incorrectly marked.",
    "False-positive and false-negative rates.",
    "Repeatability across repeated scans of the same sample.",
    "Time per scan and time required to generate the overlay.",
    "Minimum detectable test level, only after a defined reference method and controlled dilution series are established."
]: bullet(x)

doc.add_heading("9. Expected output", level=1)
doc.add_paragraph(
    "The laptop interface is expected to show the original instrument image, the UV-illuminated image, the NIR-illuminated image, and an overlay highlighting candidate regions that differ from the clean reference. A sensor value and selected illumination mode may also be displayed. The system should present a screening prompt such as 'Candidate residue detected — inspect/re-clean and follow the approved cleaning procedure.'"
)
doc.add_heading("10. Safety, limitations, and responsible use", level=1)
for x in [
    "UV safety: operate the 365 nm LED inside a closed, light-blocking enclosure; prevent direct eye/skin exposure and use appropriate protective procedures.",
    "Biological safety: do not use fresh human blood for an unsupervised project. Prefer non-hazardous surrogates or commercially prepared test materials, and follow institutional biosafety guidance.",
    "Electrical safety: use current-limiting resistors or rated drivers, check MOSFET pinouts, keep the circuit low-voltage, and verify all grounds and supply ratings before powering.",
    "Optical limitation: fluorescence and NIR contrast are not specific to blood; other residues and surface finishes can produce similar signals.",
    "Imaging limitation: an ordinary webcam may contain an IR-cut filter and may be insensitive to NIR; camera suitability must be tested.",
    "Clinical limitation: this prototype is not a sterilization monitor and cannot establish that an instrument is sterile or safe for patient use.",
    "Algorithm limitation: simple thresholding highlights image differences, not a chemically confirmed blood stain."
]: bullet(x)

doc.add_heading("11. Novelty and future patent potential", level=1)
doc.add_paragraph(
    "The general use of UV illumination, IR/NIR imaging, cameras, and image processing is not by itself enough to establish patent novelty. Any potential invention would need to be defined precisely and compared with existing patents and publications. Possible areas for further investigation include a specific low-cost optical arrangement, sequential illumination and image-fusion method, a reproducible calibration procedure, or an instrument-shaped scanning fixture. These are research directions, not claims that the concept is novel or patentable."
)
doc.add_paragraph(
    "Before public disclosure, maintain dated lab notes, circuit revisions, source-code versions, test datasets, and results. Discuss disclosure and prior-art searching with the institution's innovation/IP cell or a qualified patent professional. Do not claim a patent has been filed unless it has actually been filed."
)

doc.add_heading("12. Project schedule", level=1)
table(["Week","Task","Deliverable"],[
    ("1","Confirm component availability; review datasheets and safety plan.","Final bill of materials and approved test plan."),
    ("2","Build LED drivers and Arduino control circuit.","Bench-tested illumination switching."),
    ("3","Set up camera and light-shielded enclosure.","Repeatable image capture."),
    ("4","Implement Python/OpenCV display and overlay.","Laptop visualization demo."),
    ("5","Run baseline, positive-control feasibility, and interference tests.","Initial test dataset and metrics."),
    ("6","Improve geometry, document findings, prepare demonstration.","Final report, presentation, and prototype.")
])

doc.add_heading("13. Conclusion", level=1)
doc.add_paragraph(
    "BloodTrace is a feasible small-scale college engineering project combining optical sensing, embedded control, and computer vision. The Arduino UNO R4 Minima provides a simple way to switch illumination and stream an auxiliary sensor reading, while a laptop performs image capture and processing. The main research question is whether combining UV and NIR observations improves the visibility of residual contamination compared with a normal image. The first prototype should focus on repeatable imaging and transparent evaluation; blood-specific detection and any clinical application require further controlled validation."
)

doc.add_heading("References", level=1)
refs = [
    "Arduino. UNO R4 Minima — official hardware documentation and specifications. https://docs.arduino.cc/hardware/uno-r4-minima",
    "Arduino. UNO R4 Minima datasheet and pinout. https://docs-content.arduino.cc/resources/datasheets/ABX00080-datasheet.pdf",
    "Centers for Disease Control and Prevention (CDC). Cleaning: Guideline for Disinfection and Sterilization in Healthcare Facilities. https://www.cdc.gov/infection-control/hcp/disinfection-sterilization/cleaning.html",
    "Centers for Disease Control and Prevention (CDC). Summary of Recommendations: Disinfection and Sterilization in Healthcare Facilities. https://www.cdc.gov/infection-control/hcp/disinfection-sterilization/summary-recommendations.html",
    "OpenCV. Open Source Computer Vision Library documentation. https://docs.opencv.org/"
]
for ref in refs:
    doc.add_paragraph(ref, style="List Number")

# Footer page numbers
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("BloodTrace | College Project Proposal  •  ")
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    footer._p.append(fldChar1); footer._p.append(instrText); footer._p.append(fldChar2)

doc.save(docx_path)
print(f"Created: {docx_path}")
print(f"System diagram: {diagram_path}")
print(f"Circuit diagram: {circuit_path}")


